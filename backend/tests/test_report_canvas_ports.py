import base64,copy,io,json
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch
import pytest
from PIL import Image
from pypdf import PdfReader
from openpyxl import load_workbook
from django.core.management import call_command
from rest_framework.exceptions import ValidationError,PermissionDenied
from core.models import Unit
from operations.models import ReportDefinition
from operations.report_canvas.ports import make_layout
from operations.report_canvas.sources import resolve_report
from operations.report_canvas.engine import export,definition_snapshot
from operations.report_canvas.native import medical_tables
from operations.reports import local_report_data

@pytest.fixture
def report(setup):
 return ReportDefinition.objects.create(unit=setup["unit"],code="contacts",name="Contatos",source="local",handler="contacts",enabled=True,reviewed=True,parameter_schema={"start":"date","end":"date"},output_format="pdf")

@pytest.mark.parametrize("fmt",["pdf","xlsx"])
def test_local_port_uses_original_rows_and_exports_all_dates(report,fmt):
 original=local_report_data(report,{})
 assert original[1]
 v=make_layout(report)
 content,count=export(report,v,{},[],fmt)
 assert count==len(original[1])
 raw=base64.b64decode(content)
 if fmt=="pdf":assert "PACIENTE FICT" in " ".join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages)
 else:
  book=load_workbook(io.BytesIO(raw))
  assert any(c.value==original[1][0][0] for sheet in book for row in sheet for c in row)


def test_porting_is_idempotent_and_preserves_activation(report,tmp_path):
 enabled=report.enabled;version=report.version
 call_command("port_report_canvas",apply=True,backup=str(tmp_path/"backup.json"))
 report.refresh_from_db();assert report.canvas_layout["pages"][0]["blocks"] and not report.canvas_enabled
 assert report.enabled==enabled and report.version==version+1
 call_command("port_report_canvas",apply=True,backup=str(tmp_path/"backup2.json"))
 report.refresh_from_db();assert report.version==version+1
 assert json.loads((tmp_path/"backup.json").read_text())[0]["canvas_layout"]=={}


def test_custom_drafts_are_preserved(report,tmp_path):
 report.canvas_layout={"custom":True};report.save()
 call_command("port_report_canvas",apply=True,backup=str(tmp_path/"backup.json"))
 report.refresh_from_db();assert report.canvas_layout=={"custom":True}


def test_native_sources_cannot_cross_units_or_select_unknown_dataset(report):
 other=Unit.objects.create(slug="other",name="Outra")
 owner=ReportDefinition.objects.create(unit=other,code="owner",name="Owner")
 with pytest.raises(ValidationError):resolve_report(owner,{"report_id":report.pk})
 with pytest.raises(ValidationError):resolve_report(report,{"report_id":report.pk,"dataset":"unknown"})
 report.enabled=False;report.save()
 assert resolve_report(report,{"report_id":report.pk}).pk==report.pk
 foreign=ReportDefinition.objects.create(unit=report.unit,code="foreign",name="Foreign")
 with pytest.raises(ValidationError):resolve_report(foreign,{"report_id":report.pk})


def test_native_source_snapshot_is_sensitive_to_handler_change(report):
 report.canvas_layout=make_layout(report);report.save()
 a=definition_snapshot(report)
 report.handler="confirmed";report.save()
 assert definition_snapshot(report)!=a

@pytest.fixture
def medical_data(tmp_path):
 layout={"TBRELATORIO_CORES":[{"ID_RELATORIO":63,"CRITERIO":"CLINICA GERAL"}]}
 (tmp_path/"layout.json").write_text(json.dumps(layout))
 for name in ["first-background.jpg","last-background.jpg","logo.png"]:Image.new("RGB",(100,60),"white").save(tmp_path/name)
 d={"clock":[{"DB_TIMESTAMP":"2026-10-23T00:00:00"}],"generated_at":"2026-10-25T02:00:00-03:00", "freshness":{k:[{"ULTIMA_ALTA" if k=="internacao" else "ULTIMA_DATA":"2026-10-22"}] for k in ["internacao","exames","atendimentos","parecer"]},
 "stays":[{"ANO":2026,"MES":10,"TIPO_CONVENIO":"SUS","DS_ESPECIALID":"CLINICA GERAL","DIAS":20,"QUANTIDADE":4}],
 "exams":[{"ANO":2026,"MES":10,"QUANTIDADE":8}],"attendances":[{"ANO":2026,"MES":10,"QUANTIDADE":4}],
 "exams_days":[{"DATA":"2026-10-22","QUANTIDADE":8}],"attendances_days":[{"DATA":"2026-10-22","QUANTIDADE":4}],
 "opinions_months":[{"ANO":2026,"MES":10,"DS_ESPECIALID":"CLINICA GERAL","DIAS":"0.5","QUANTIDADE":2}],
 "opinions_days":[{"DATA":"2026-10-22","DS_ESPECIALID":"CLINICA GERAL","DIAS":"0.5","QUANTIDADE":2}]}
 return tmp_path,d


def test_medical_calculations_preserve_windows_weighted_averages_and_times(medical_data):
 root,d=medical_data
 with patch("operations.medical_indicators.asset_root",return_value=root):tables=medical_tables(d)
 cols,rows=tables["stays"]
 assert len(rows)==15 and rows[1][1]==5 and rows[-1][1]==5
 assert tables["ratios"][1][1]==["2026",2.0,2.5]
 assert len(tables["daily"][1])==8
 assert len(tables["opinions_days"][0])==11
 assert any("06:00:00" in row for row in tables["opinions_months"][1])
 assert tables["metadata"][1][0][0]=="2026-10-22"

@pytest.mark.parametrize("fmt",["pdf","xlsx"])
def test_medical_layout_is_editable_and_collects_only_once(report,medical_data,fmt):
 root,d=medical_data
 report.source="composite";report.handler="medical_indicators_v1";report.parameter_schema={"data_referencia":{"type":"date","required":False}};report.save()
 with patch("operations.medical_indicators.asset_root",return_value=root),patch("operations.medical_indicators.collect_data",return_value=d) as collect:
  v=make_layout(report)
  assert len(v["pages"])==7 and len(v["sources"])==8
  assert any(b["type"]=="line" for p in v["pages"] for b in p["blocks"])
  content,count=export(report,v,{"data_referencia":"2026-10-23"},[],fmt)
  assert collect.call_count==1
  raw=base64.b64decode(content)
  if fmt=="pdf":
   pages=PdfReader(io.BytesIO(raw)).pages
   assert len(pages)>=7
   assert "22/10/2026" in pages[1].extract_text()
  else:assert len(load_workbook(io.BytesIO(raw)).worksheets)>=7


def test_api_preserves_native_dataset_and_lists_allowed_sources(setup,report):
 url=f"/api/v1/reports/{report.pk}/canvas/?unit=ame"
 response=setup["client"].get(url)
 assert response.status_code==200
 src=next(s for s in response.data["sources"] if s["id"]==report.pk)
 assert src["datasets"][0]["id"]=="rows"


@pytest.mark.parametrize("handler",["medical_indicators_v1","stock_loans_v1"])
def test_port_does_not_bypass_automatic_send_restrictions(report,setup,handler):
 from operations.reporting import make_execution
 report.canvas_enabled=True;report.handler=handler;report.save()
 with pytest.raises(PermissionDenied):make_execution(report,{},user=setup["user"],origin="automatic")


def test_local_canvas_preserves_maximum_period(report):
 with pytest.raises(ValidationError):export(report,make_layout(report),{"start":"2020-01-01","end":"2026-01-01"},[],"pdf")



def test_widths_are_validated_and_used(report):
 from operations.report_canvas.schema import validate_layout
 v=make_layout(report);b=v["pages"][0]["blocks"][1]
 b["options"]["column_widths"]=[200]+[70]*7
 content,_=export(report,v,{},[],"pdf");assert base64.b64decode(content).startswith(b"%PDF")
 b["options"]["column_widths"]=[0]
 with pytest.raises(ValidationError):validate_layout(v)
 b["options"]["column_widths"]=[1]
 with pytest.raises(ValidationError):export(report,v,{},[],"pdf")


def test_missing_native_dataset_cannot_overwrite_ported_layout(report,setup):
 import uuid
 report.handler="medical_indicators_v1";report.save()
 v={"schema_version":1,"page":{"width":794,"height":1123,"columns":24,"background":"#ffffff"},"sources":[{"id":"data","type":"report","report_id":report.pk}],"pages":[{"id":"p1","name":"Página","blocks":[]}]}
 response=setup["client"].post(f"/api/v1/reports/{report.pk}/canvas/?unit=ame",{"version":report.version,"layout":v,"enabled":False},format="json",HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()))
 assert response.status_code==400
 report.refresh_from_db();assert report.canvas_layout=={}


def test_activation_switches_only_generator_and_is_idempotent(report,tmp_path):
 from operations.models import ReportGrant,ReportAutomation
 report.parameter_schema={};report.save()
 before=(report.enabled,report.parameter_schema,report.handler,ReportGrant.objects.count(),ReportAutomation.objects.count())
 with patch("operations.management.commands.activate_report_canvas.export",return_value=(base64.b64encode(b"%PDF-test").decode(),3)):
  call_command("activate_report_canvas",apply=True,backup=str(tmp_path/"before.json"))
 report.refresh_from_db()
 assert report.canvas_enabled and report.canvas_layout
 assert (report.enabled,report.parameter_schema,report.handler,ReportGrant.objects.count(),ReportAutomation.objects.count())==before
 version=report.version
 call_command("activate_report_canvas",apply=True,backup=str(tmp_path/"again.json"))
 report.refresh_from_db();assert report.version==version


def test_activation_failure_keeps_original_generator(report,tmp_path):
 from django.core.management.base import CommandError
 with patch("operations.management.commands.activate_report_canvas.export",side_effect=ValidationError("Falha")):
  with pytest.raises(CommandError):call_command("activate_report_canvas",apply=True,backup=str(tmp_path/"before.json"))
 report.refresh_from_db();assert not report.canvas_enabled and report.canvas_layout=={}


def test_beds_canvas_xlsx_keeps_exact_four_columns_without_title_rows(report):
 from operations.report_canvas.engine import export
 from operations.report_canvas.sources import Dataset
 report.handler="bed_snapshots_v1";report.output_format="xlsx";report.save()
 layout=make_layout(report)
 records=[[1,"Clínica",10,date(2026,10,1)],[2,"UTI",4,date(2026,10,1)],[1,"Clínica",11,date(2026,10,2)]]
 with patch("operations.report_canvas.native.read",return_value=Dataset(["CD_UNID_INT","DS_UNID_INT","QUANTIDADE","DATA"],records)):
  content,count=export(report,layout,{},[],"xlsx")
 book=load_workbook(io.BytesIO(base64.b64decode(content)))
 assert len(book.worksheets)==1
 sheet=book.active
 assert sheet.max_column==4 and sheet.max_row==4 and count==3
 assert list(next(sheet.values))==["CD_UNID_INT","DS_UNID_INT","QUANTIDADE","DATA"]
 assert sheet["C4"].value==11
 assert sheet["A2"].fill.fgColor.rgb==sheet["A3"].fill.fgColor.rgb
 assert sheet["A3"].fill.fgColor.rgb!=sheet["A4"].fill.fgColor.rgb


def test_stock_canvas_keeps_subtotals_incomplete_values_and_one_query(report):
 from operations.report_canvas.native import read
 report.handler="stock_loans_v1";report.source="sql";report.save()
 ctx=SimpleNamespace(native_cache={})
 cols=["DS_FORNECEDOR","TP_EMPRESTIMO","DT_EMPRESTIMO","CPRODUTO","PRODUTO","UNIDADE","QT_ENTRADA","QT_ESTOQUE_ATUAL","DS_LOCALIZACAO"]
 rows=[["A","E",date(2026,10,1),1,"Produto","UN",3,5,"Local"],["A","E",date(2026,10,1),2,"Outro","UN",4,None,"Local"]]
 with patch("operations.stock_loans.prepare_values",return_value={}),patch("operations.report_engine.run_query",return_value=(cols,rows)) as query:
  details=read(ctx,report,"rows",{})
  totals=read(ctx,report,"summary",{})
 assert query.call_count==1 and len(details.rows)==2
 assert totals.rows[0]==["A",2,7,"Incompleto"]
 assert totals.rows[-1]==["Total geral",2,7,"Incompleto"]
 assert len(make_layout(report)["pages"])==2
