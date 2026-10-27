import base64
import io
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

import pytest
from django.db import close_old_connections, connections
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from core.models import Membership, Unit
from operations.bed_snapshots import (HANDLER, HEADERS, collect_snapshot, definition,
    export_selection, queue_collection)
from operations.models import (BedSnapshotRow, ReportConnection, ReportDefinition,
    ReportSnapshotConfiguration, ReportSnapshotRun)
from operations.reporting import make_execution, execute_report, readiness
from operations.report_engine import report_filename, render

NOW = datetime(2026,10,24,4,21,tzinfo=ZoneInfo("America/Sao_Paulo"))


@pytest.fixture
def beds(setup):
    unit=setup["unit"]
    conn=ReportConnection.objects.create(unit=unit,name="Oracle HML",engine="oracle",host="hml.invalid",database="hml",enabled=True,validated=True)
    source=ReportDefinition.objects.create(unit=unit,code="bed-source",name="Source",catalog_kind="component",source="sql",connection=conn,
        query_sql="SELECT CD_UNID_INT, DS_UNID_INT, QUANTIDADE, DATA FROM dbamv.leito",reviewed=True,enabled=True,parameter_schema={})
    report=ReportDefinition.objects.create(unit=unit,code="bed-report",name="Tabela Leitos",source="local",handler=HANDLER,reviewed=True,enabled=True,
        output_format="xlsx",parameter_schema={"WPA_DATA_INICIAL":{"type":"date","required":True},"WPA_DATA_FINAL":{"type":"date","required":True}})
    config=ReportSnapshotConfiguration.objects.create(unit=unit,code="beds",name="Beds",source_report=source,output_report=report,
        environment="homologation",collection_time=time(4,20,23),enabled=True)
    with patch("operations.bed_snapshots.timezone.now",return_value=NOW):
        yield {**setup,"conn":conn,"source":source,"report":report,"config":config}


def fake_rows(quantity="41", name="ALA B"):
    return HEADERS, [["2",name,quantity,"2026-10-24 04:21:00"],["3","ALA C","15","2026-10-24 04:21:00"]]


def collect(env, result=None):
    run=queue_collection(env["config"].pk,user=env["user"],dispatch=False)
    with patch("operations.bed_snapshots.run_query",return_value=result or fake_rows()) as query:
        collect_snapshot(str(run.pk))
    run.refresh_from_db()
    return run,query


def period(start="2026-10-24",end="2026-10-24"):
    return {"WPA_DATA_INICIAL":start,"WPA_DATA_FINAL":end}


def test_collect_stores_source_clock_and_deduplicates_without_transport(beds):
    with patch("operations.report_whatsapp.service") as sender:
        run,query=collect(beds)
        duplicate=queue_collection(beds["config"].pk,user=beds["user"],dispatch=False)
        collect_snapshot(str(duplicate.pk))
    assert run.state=="ready" and run.row_count==2 and run.reference_date==NOW.date()
    assert run.source_at==NOW and run.captured_at==NOW
    assert duplicate.pk==run.pk and ReportSnapshotRun.objects.count()==1
    assert BedSnapshotRow.objects.filter(run=run).count()==2
    query.assert_called_once_with(beds["source"],{})
    sender.assert_not_called()


@pytest.mark.parametrize("result",[
    (HEADERS,[]),
    (HEADERS,[["2","ALA","-1","2026-10-24 04:21:00"]]),
    (HEADERS,[["2","ALA","1.5","2026-10-24 04:21:00"]]),
    (HEADERS,[["2","ALA","1","2026-10-23 04:21:00"]]),
    (HEADERS,[["2","ALA","1","2026-10-24 04:21:00"]]*2),
    (["NOME"],[["ALA"]]),
])
def test_invalid_results_never_leave_partial_snapshots(beds,result):
    run,_=collect(beds,result)
    assert run.state=="failed" and run.error_code=="invalid_result"
    assert not BedSnapshotRow.objects.exists()


def test_failure_redacts_driver_details_and_can_retry(beds):
    run=queue_collection(beds["config"].pk,dispatch=False)
    with patch("operations.bed_snapshots.run_query",side_effect=RuntimeError("password=never-log-this")):
        collect_snapshot(str(run.pk))
    run.refresh_from_db()
    assert run.state=="failed" and run.error_code=="source_query_failed"
    retry=queue_collection(beds["config"].pk,user=beds["user"],dispatch=False)
    assert retry.pk==run.pk
    with patch("operations.bed_snapshots.run_query",return_value=fake_rows()):collect_snapshot(str(run.pk))
    run.refresh_from_db();assert run.state=="ready" and run.attempt_count==2


def test_source_change_during_collection_discards_result(beds):
    run=queue_collection(beds["config"].pk,dispatch=False)
    def change(*args,**kwargs):
        ReportSnapshotConfiguration.objects.filter(pk=beds["config"].pk).update(version=2)
        return fake_rows()
    with patch("operations.bed_snapshots.run_query",side_effect=change):collect_snapshot(str(run.pk))
    run.refresh_from_db()
    assert run.state=="failed" and run.error_code=="source_changed"
    assert not BedSnapshotRow.objects.exists()


def test_revoked_requester_cannot_finish_collection(beds):
    run=queue_collection(beds["config"].pk,user=beds["user"],dispatch=False)
    def revoke(*args,**kwargs):
        type(beds["user"]).objects.filter(pk=beds["user"].pk).update(is_active=False)
        return fake_rows()
    with patch("operations.bed_snapshots.run_query",side_effect=revoke):collect_snapshot(str(run.pk))
    run.refresh_from_db();assert run.state=="failed" and not BedSnapshotRow.objects.exists()


def test_scheduler_does_not_run_early_or_backfill_outside_window(beds):
    cfg=beds["config"]
    assert queue_collection(cfg.pk,origin="scheduled",now=NOW.replace(hour=4,minute=19),dispatch=False) is None
    run=queue_collection(cfg.pk,origin="scheduled",now=NOW.replace(hour=7),dispatch=False)
    assert run.state=="missed" and run.error_code=="collection_window_missed"
    assert run.reference_date==NOW.date() and not BedSnapshotRow.objects.exists()
    cfg.enabled=False;cfg.save()
    assert queue_collection(cfg.pk,origin="scheduled",dispatch=False) is None


def test_weekdays_and_source_environment_isolate_datasets(beds):
    cfg=beds["config"];cfg.weekdays=[0,1,2,3,4];cfg.save()
    assert queue_collection(cfg.pk,origin="scheduled",dispatch=False) is None  # Saturday
    run,_=collect(beds)
    assert export_selection(beds["report"],period())[1][0].pk==run.pk
    old_key=run.dataset_key
    cfg.environment="production";cfg.save()
    assert definition(cfg)["dataset_key"]!=old_key
    with pytest.raises(ValidationError,match="snapshots"):export_selection(beds["report"],period())


def test_xlsx_matches_legacy_columns_native_types_dates_and_prevents_formulas(beds):
    run,_=collect(beds,fake_rows(name="=SUM(1,2)"))
    assert readiness(beds["report"])==[]
    with patch("operations.report_whatsapp.service") as sender:
        job=make_execution(beds["report"],period(),user=beds["user"])
        execute_report(str(job.pk))
    job.refresh_from_db();assert job.state=="ready" and job.format=="xlsx" and job.row_count==2
    book=load_workbook(io.BytesIO(base64.b64decode(job.result)))
    assert book.sheetnames==["LEITOS_SISTEMA"]
    sheet=book.active
    assert sheet.max_column==4 and list(next(sheet.values))==HEADERS
    assert sheet.cell(2,1).value==2 and sheet.cell(2,3).value==41
    assert sheet.cell(2,2).value=="=SUM(1,2)" and sheet.cell(2,2).data_type=="s"
    assert sheet.cell(2,4).value==datetime(2026,10,24)
    assert job.snapshot["snapshot_runs"][0]["id"]==str(run.pk)
    assert job.snapshot["data_environment"]=="homologation"
    assert report_filename("Leitos","xlsx")=="Leitos.xlsx"
    sender.assert_not_called()
    response=beds["client"].get(f"/api/v1/report-executions/{job.pk}/download/?unit=ame")
    assert response.status_code==200 and response.content.startswith(b"PK")
    assert "spreadsheetml" in response["Content-Type"]


def test_missing_days_block_export_instead_of_faking_history(beds):
    collect(beds)
    with pytest.raises(ValidationError,match="23/10/2026"):
        make_execution(beds["report"],period("2026-10-23"),user=beds["user"])
    with pytest.raises(ValidationError):export_selection(beds["report"],period("2026-10-25","2026-10-24"))
    with pytest.raises(ValidationError):export_selection(beds["report"],period("2026-10-24","2026-10-25"))


def test_configuration_and_runs_api_enforce_unit_role_version_and_idempotency(beds):
    c=beds["client"];cfg=beds["config"]
    other=Unit.objects.create(slug="other",name="Other")
    Membership.objects.create(unit=other,user=beds["user"],role="admin")
    assert c.get(f"/api/v1/report-snapshots/{cfg.pk}/?unit=other").status_code==404
    assert c.post(f"/api/v1/report-snapshots/{cfg.pk}/collect/?unit=ame",{"version":999},format="json",HTTP_IDEMPOTENCY_KEY="bed-wrong-version").status_code==409
    data={"version":cfg.version}
    with patch("operations.bed_snapshots.enqueue_collection"):
        first=c.post(f"/api/v1/report-snapshots/{cfg.pk}/collect/?unit=ame",data,format="json",HTTP_IDEMPOTENCY_KEY="bed-collect-1")
        second=c.post(f"/api/v1/report-snapshots/{cfg.pk}/collect/?unit=ame",data,format="json",HTTP_IDEMPOTENCY_KEY="bed-collect-1")
    assert first.status_code==200,first.data
    assert second.data==first.data and ReportSnapshotRun.objects.count()==1
    run=ReportSnapshotRun.objects.get()
    assert c.get(f"/api/v1/report-snapshot-runs/{run.pk}/records/?unit=other").status_code==404
    member=Membership.objects.get(unit=beds["unit"],user=beds["user"]);member.role="operator";member.save()
    assert c.get("/api/v1/report-snapshots/?unit=ame").status_code==403
    assert c.get("/api/v1/report-snapshot-runs/?unit=ame").status_code==403


@pytest.mark.django_db(transaction=True)
def test_two_workers_queue_and_collect_only_one_daily_snapshot(beds):
    cfg=beds["config"]
    def queue():
        close_old_connections()
        try:return queue_collection(cfg.pk,dispatch=False).pk
        finally:connections.close_all()
    with ThreadPoolExecutor(max_workers=2) as pool:ids=list(pool.map(lambda _:queue(),range(2)))
    assert ids[0]==ids[1] and ReportSnapshotRun.objects.count()==1
    def worker():
        close_old_connections()
        try:collect_snapshot(str(ids[0]))
        finally:connections.close_all()
    with patch("operations.bed_snapshots.run_query",return_value=fake_rows()) as query:
        with ThreadPoolExecutor(max_workers=2) as pool:list(pool.map(lambda _:worker(),range(2)))
    assert query.call_count==1
    assert ReportSnapshotRun.objects.get().state=="ready" and BedSnapshotRow.objects.count()==2


def test_generic_xlsx_also_preserves_literal_strings():
    payload=render("Example",["A"],[["=1+1"]],"xlsx")
    sheet=load_workbook(io.BytesIO(base64.b64decode(payload))).active
    assert sheet.cell(2,1).value=="=1+1" and sheet.cell(2,1).data_type=="s"


# Real legacy history is a separate series; missing dates and unknown clocks stay explicit.
def legacy_payload():
    return {"source_host":"172.16.0.123", "source_database":"C:/WhatsReport/Servidor/BD/GLDADOS.FDB",
        "source_table":"LEITOS_SISTEMA", "environment":"production", "captured_time_known":False,
        "start":"2026-10-21", "end":"2026-10-23", "rows":[
            {"CD_UNID_INT":code,"DS_UNID_INT":name,"QUANTIDADE":quantity,"DATA":day+" 00:00:00"}
            for day in ("2026-10-22","2026-10-23") for code,name,quantity in ((2,"ALA B",41),(3,"ALA C",15))]}


def test_legacy_import_preserves_gaps_unknown_clocks_and_is_idempotent(beds):
    from operations.legacy_bed_snapshots import import_history
    with patch("operations.bed_snapshots.run_query") as query,patch("operations.report_whatsapp.service") as send:
        first=import_history(beds["config"],legacy_payload(),checksum="a"*64,actor=beds["user"])
        again=import_history(beds["config"],legacy_payload(),checksum="a"*64,actor=beds["user"])
    assert first["imported_days"]==2 and first["missing_dates"]==["2026-10-21"]
    assert again["imported_days"]==0 and again["unchanged_days"]==2
    assert ReportSnapshotRun.objects.count()==3 and BedSnapshotRow.objects.count()==4
    for run in ReportSnapshotRun.objects.all():
        assert run.origin=="legacy" and run.captured_at is None and run.source_at is None and run.scheduled_for is None
        assert run.source_snapshot["environment"]=="production" and not run.source_snapshot["captured_time_known"]
    missing=ReportSnapshotRun.objects.get(reference_date="2026-10-21")
    assert missing.state=="missed" and missing.error_code=="legacy_snapshot_unavailable" and not missing.rows.exists()
    query.assert_not_called();send.assert_not_called()


def test_legacy_export_uses_only_real_rows_and_identifies_missing_dates(beds):
    from operations.legacy_bed_snapshots import import_history
    import_history(beds["config"],legacy_payload(),checksum="a"*64,actor=beds["user"])
    with patch("operations.report_whatsapp.service") as send:
        job=make_execution(beds["report"],period("2026-10-21","2026-10-23"),user=beds["user"])
        assert job.snapshot["bed_history"]["origin"]=="legacy"
        execute_report(str(job.pk))
    job.refresh_from_db();assert job.state=="ready" and job.row_count==4
    assert job.snapshot["data_environment"]=="production" and job.snapshot["data_origin"]=="legacy"
    assert job.snapshot["missing_dates"]==["2026-10-21"]
    assert all(s["captured_at"] is None and s["imported_at"] for s in job.snapshot["snapshot_runs"])
    sheet=load_workbook(io.BytesIO(base64.b64decode(job.result))).active
    assert sheet.max_row==5 and {r[3].date().isoformat() for r in list(sheet.values)[1:]}=={"2026-10-22","2026-10-23"}
    response=beds["client"].get(f"/api/v1/report-executions/{job.pk}/download/?unit=ame")
    assert response.status_code==200 and "Legado" in response["Content-Disposition"]
    response=beds["client"].get(f"/api/v1/report-executions/{job.pk}/?unit=ame")
    assert response.data["missing_dates"]==["2026-10-21"] and response.data["data_origin"]=="legacy"
    send.assert_not_called()


def test_homologation_and_legacy_never_merge_in_the_same_export(beds):
    from operations.legacy_bed_snapshots import import_history
    import_history(beds["config"],legacy_payload(),checksum="a"*64)
    collect(beds)
    with pytest.raises(ValidationError,match="homologação"):
        make_execution(beds["report"],period("2026-10-22","2026-10-24"),user=beds["user"])
    assert export_selection(beds["report"],period())[1][0].origin=="manual"
    assert export_selection(beds["report"],period("2026-10-22","2026-10-23"))[1][0].origin=="legacy"


def test_conflicting_reimport_rolls_back_the_whole_file(beds):
    from operations.legacy_bed_snapshots import import_history
    import_history(beds["config"],legacy_payload(),checksum="a"*64)
    changed=legacy_payload();changed["start"]="2026-10-20";changed["rows"][0]["QUANTIDADE"]=99
    changed["rows"].append({"CD_UNID_INT":2,"DS_UNID_INT":"ALA B","QUANTIDADE":10,"DATA":"2026-10-20 00:00:00"})
    with pytest.raises(ValidationError,match="conflita"):import_history(beds["config"],changed,checksum="b"*64)
    assert ReportSnapshotRun.objects.count()==3 and not ReportSnapshotRun.objects.filter(reference_date="2026-10-20").exists()
    assert BedSnapshotRow.objects.filter(ward_code=2,quantity=41).count()==2


@pytest.mark.parametrize("change",[{"environment":"homologation"},{"captured_time_known":True},{"end":"2026-10-24"}])
def test_legacy_import_rejects_wrong_environment_or_false_capture_information(beds,change):
    from operations.legacy_bed_snapshots import import_history
    payload={**legacy_payload(),**change}
    with pytest.raises(ValidationError):import_history(beds["config"],payload,checksum="a"*64)
    assert not ReportSnapshotRun.objects.exists()


def test_historical_idempotency_uses_reference_date_across_timezones(beds):
    from operations.legacy_bed_snapshots import import_history
    beds["unit"].timezone="Asia/Tokyo";beds["unit"].save()
    first=import_history(beds["config"],legacy_payload(),checksum="a"*64)
    second=import_history(beds["config"],legacy_payload(),checksum="a"*64)
    assert first["imported_days"]==2 and second["unchanged_days"]==2



def test_delivery_feedback_explicitly_identifies_real_history_gaps():
    from types import SimpleNamespace
    from operations.report_whatsapp import report_delivery_feedback
    job=SimpleNamespace(delivery_state="sent_baileys",snapshot={"data_origin":"legacy","missing_dates":["2026-10-04","2026-10-10"]})
    message=report_delivery_feedback(job)
    assert "produção" in message and "04/10/2026" in message and "10/10/2026" in message
    job.snapshot={}
    assert report_delivery_feedback(job)=="Relatório enviado."
