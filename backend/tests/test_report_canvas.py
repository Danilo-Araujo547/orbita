import base64
import copy
import io
import sqlite3
import uuid
from unittest.mock import patch

import pytest
from django.contrib.auth import get_user_model
from rest_framework.exceptions import PermissionDenied, ValidationError
from pypdf import PdfReader
from openpyxl import load_workbook
from PIL import Image
from core.models import Membership, Unit
from operations.models import ReportConnection, ReportDefinition, ReportExecution
from operations.report_canvas.schema import default_layout, validate_layout
from operations.report_canvas.engine import export, definition_snapshot, check_snapshot
from operations.reporting import make_execution, execute_report


def layout():
    v=default_layout()
    v["sources"]=[{"id":"manual","type":"inline","columns":["Mes","Total"],"rows":[["Jan",10],["Fev",20]]}]
    v["pages"][0]["blocks"]=[{"id":"title","type":"text","x":40,"y":25,"width":700,"height":60,"text":"Relatório {{report_name}}"},{"id":"total","type":"kpi","x":40,"y":100,"width":200,"height":100,"source_id":"manual","title":"Total","options":{"value_field":"Total","aggregate":"sum"}},{"id":"chart","type":"bar","x":40,"y":230,"width":650,"height":260,"source_id":"manual","title":"Mensal","options":{"x_field":"Mes","y_fields":["Total"]}},{"id":"table","type":"table","x":40,"y":540,"width":650,"height":280,"source_id":"manual","title":"Registros"}]
    return v


@pytest.fixture
def env(setup,tmp_path,settings):
    settings.ASYNC_TASKS=False
    settings.REPORT_SQLITE_ROOT=str(tmp_path)
    report=ReportDefinition.objects.create(unit=setup["unit"],code="canvas-test",name="Canvas Teste",source="legacy_pending",enabled=True,reviewed=True,canvas_enabled=True,canvas_layout=layout(),monthly_goals=[{"year":2026,"month":10,"value":19420}])
    path=tmp_path/"source.sqlite3"
    db=sqlite3.connect(path);db.execute("CREATE TABLE totals (label TEXT,total INTEGER)");db.executemany("INSERT INTO totals VALUES (?,?)",[("Jan",10),("Fev",20)]);db.commit();db.close()
    conn=ReportConnection.objects.create(unit=setup["unit"],name="Canvas source",engine="sqlite",database=str(path),enabled=True,validated=True)
    child=ReportDefinition.objects.create(unit=setup["unit"],code="canvas-sql",name="Dados SQL",source="sql",query_sql="SELECT label,total FROM totals",connection=conn,enabled=True,reviewed=True,catalog_kind="component")
    return setup,report,child


def post(e,r,action,data):
    return e["client"].post(f"/api/v1/reports/{r.pk}/{action}/?unit=ame",data,format="json",HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()))


def test_pdf_xlsx_native_charts_and_untrusted_cells(env):
    e,r,_=env
    r.canvas_layout["sources"][0]["rows"][0][0]="=HYPERLINK(\"bad\")"
    content,count=export(r,r.canvas_layout,{},r.monthly_goals,"pdf")
    pdf=PdfReader(io.BytesIO(base64.b64decode(content)))
    assert count==2 and len(pdf.pages)==1
    assert "30" in pdf.pages[0].extract_text() and "Canvas Teste" in pdf.pages[0].extract_text()
    content,count=export(r,r.canvas_layout,{},r.monthly_goals,"xlsx")
    book=load_workbook(io.BytesIO(base64.b64decode(content)))
    assert sum(len(s._charts) for s in book)==1
    strings=[c for s in book for row in s for c in row if isinstance(c.value,str) and c.value.startswith("=HYPERLINK")]
    assert strings and all(c.data_type=="s" for c in strings)


@pytest.mark.parametrize("kind",["line","pie"])
def test_chart_strategies(env,kind):
    _,r,_=env
    r.canvas_layout["pages"][0]["blocks"][2]["type"]=kind
    for fmt in ("pdf","xlsx"):
        content,_=export(r,r.canvas_layout,{},[],fmt)
        assert len(base64.b64decode(content))>500


def test_large_table_continues_without_losing_rows(env):
    _,r,_=env
    v=layout();v["sources"][0]["rows"]=[[str(i),i] for i in range(90)]
    v["pages"][0]["blocks"]=[{**v["pages"][0]["blocks"][-1],"height":160,"y":40}]
    content,_=export(r,v,{},[],"pdf");pdf=PdfReader(io.BytesIO(base64.b64decode(content)))
    assert len(pdf.pages)>1 and "89" in pdf.pages[-1].extract_text()
    content,_=export(r,v,{},[],"xlsx");book=load_workbook(io.BytesIO(base64.b64decode(content)))
    assert any(s.max_row==91 for s in book)


@pytest.mark.parametrize("change",["bounds","nan","duplicate","missing_source","huge","unknown"])
def test_layout_validation(change):
    v=layout()
    if change=="bounds":v["pages"][0]["blocks"][0]["x"]=9999
    if change=="nan":v["page"]["width"]=float("nan")
    if change=="duplicate":v["pages"].append(copy.deepcopy(v["pages"][0]))
    if change=="missing_source":v["pages"][0]["blocks"][1]["source_id"]="unknown"
    if change=="huge":v["sources"][0]["rows"]=[["x",1]]*2001
    if change=="unknown":v["pages"][0]["blocks"][0]["type"]="javascript"
    with pytest.raises(ValidationError):validate_layout(v)


def test_admin_config_conflict_draft_preserves_handler_and_unit_isolation(env):
    e,r,child=env;r.canvas_enabled=False;r.handler="medical_indicators_v1";r.enabled=False;r.save()
    data={"version":r.version,"layout":layout(),"enabled":False,"output_format":"xlsx"}
    response=post(e,r,"canvas",data);assert response.status_code==200,response.data
    r.refresh_from_db();assert not r.canvas_enabled and r.handler=="medical_indicators_v1" and r.output_format=="pdf"
    assert post(e,r,"canvas",data).status_code==409
    other=Unit.objects.create(slug="other",name="Other")
    foreign=ReportDefinition.objects.create(unit=other,code="foreign",name="Foreign",source="sql")
    data["version"]=r.version;data["layout"]["sources"]=[{"id":"manual","type":"report","report_id":foreign.pk}]
    assert post(e,r,"canvas",data).status_code==400
    user=get_user_model().objects.create_user("viewer-canvas");Membership.objects.create(unit=e["unit"],user=user,role="viewer");e["client"].force_authenticate(user)
    assert post(e,r,"canvas-preview",{"layout":layout()}).status_code==403


def test_preview_has_no_jobs_or_message_sends(env):
    e,r,_=env
    before=ReportExecution.objects.count()
    response=post(e,r,"canvas-preview",{"layout":layout(),"parameters":{},"output_format":"pdf"})
    assert response.status_code==200,response.data
    assert response.data["rows"]==2 and ReportExecution.objects.count()==before


def test_sql_goals_snapshots_and_change_rejection(env):
    e,r,child=env
    v=layout();v["sources"]=[{"id":"manual","type":"report","report_id":child.pk}];v["pages"][0]["blocks"][1]["options"]["value_field"]="total";v["pages"][0]["blocks"][2]["options"]={"x_field":"label","y_fields":["total"]}
    r.canvas_layout=v;r.save()
    job=make_execution(r,{},user=e["user"])
    execute_report(job.pk);job.refresh_from_db();assert job.state=="ready",job.error_code
    assert job.row_count==2
    child.version+=1;child.save()
    with pytest.raises(PermissionDenied):check_snapshot(job)
    v=layout();v["sources"]=[{"id":"manual","type":"goals"}];v["pages"][0]["blocks"]=[{**v["pages"][0]["blocks"][1],"options":{"value_field":"value"}}]
    content,count=export(r,v,{},r.monthly_goals,"pdf")
    assert count==1 and "19.420" in PdfReader(io.BytesIO(base64.b64decode(content))).pages[0].extract_text()


def test_asset_ownership_and_render(env):
    e,r,_=env
    out=io.BytesIO();Image.new("RGBA",(20,20),(20,150,200,200)).save(out,format="PNG")
    body={"filename":"logo.png","content":base64.b64encode(out.getvalue()).decode()}
    response=post(e,r,"canvas-assets",body);assert response.status_code==200,response.data
    a=response.data
    v=layout();v["pages"][0]["blocks"].append({"id":"logo","type":"image","x":600,"y":100,"width":100,"height":100,"asset_id":a["id"]})
    for fmt in ("pdf","xlsx"):assert export(r,v,{},[],fmt)[0]
    configured=post(e,r,"canvas",{"version":r.version,"layout":v,"enabled":True,"output_format":"pdf"})
    assert configured.status_code==200,configured.data
    assert configured.data["assets"][0]["id"]==a["id"]
    other=ReportDefinition.objects.create(unit=r.unit,name="Other",code="other")
    with pytest.raises(ValidationError):definition_snapshot(other,v)
    assert post(e,r,"canvas-assets",{"filename":"x.svg","content":base64.b64encode(b"<svg/>").decode()}).status_code==400


def test_enable_disable_restores_original_generator_configuration(env):
    e,r,_=env
    r.canvas_enabled=False;r.output_format="csv";r.reviewed=False;r.save()
    result=post(e,r,"canvas",{"version":r.version,"layout":layout(),"enabled":True,"output_format":"xlsx"})
    assert result.status_code==200,result.data
    r.refresh_from_db();assert r.canvas_enabled and r.output_format=="xlsx" and r.reviewed
    result=post(e,r,"canvas",{"version":r.version,"layout":layout(),"enabled":False,"output_format":"xlsx"})
    assert result.status_code==200,result.data
    r.refresh_from_db();assert not r.canvas_enabled and r.output_format=="csv" and not r.reviewed


def test_source_parameter_alias_and_shared_readonly_session(env):
    e,r,child=env
    child.query_sql="SELECT label,total FROM totals WHERE total >= :minimum";child.parameter_schema={"minimum":{"type":"integer"}};child.save()
    v=layout();v["sources"]=[{"id":"manual","type":"report","report_id":child.pk,"parameters":{"minimum":{"parameter":"filter"}}}]
    v["pages"][0]["blocks"]=[{**v["pages"][0]["blocks"][1],"options":{"value_field":"total"}}]
    content,count=export(r,v,{"filter":15},[],"pdf")
    assert count==1 and "20" in PdfReader(io.BytesIO(base64.b64decode(content))).pages[0].extract_text()


def test_changed_connection_and_unauthorized_assets_block_execution(env):
    e,r,child=env
    v=layout();v["sources"]=[{"id":"manual","type":"report","report_id":child.pk}];v["pages"][0]["blocks"]=[{**v["pages"][0]["blocks"][1],"options":{"value_field":"total"}}]
    r.canvas_layout=v;r.save();job=make_execution(r,{},user=e["user"])
    child.connection.version+=1;child.connection.save()
    with pytest.raises(PermissionDenied):check_snapshot(job)
    result=e["client"].get(f"/api/v1/reports/{r.pk}/canvas-assets/?unit=ame&asset_id=not-a-uuid")
    assert result.status_code==400


def test_sql_tables_preserve_numeric_cells(env):
    _,r,child=env
    v=layout();v["sources"]=[{"id":"manual","type":"report","report_id":child.pk}]
    v["pages"][0]["blocks"]=[{**v["pages"][0]["blocks"][-1],"y":40}]
    content,_=export(r,v,{},[],"xlsx")
    book=load_workbook(io.BytesIO(base64.b64decode(content)))
    assert any(c.value==10 and c.data_type=="n" for row in book.worksheets[0] for c in row)


def test_date_placeholder_and_datetime_display(env):
    from datetime import date
    from operations.report_canvas.formatting import display
    _,r,_=env
    v=default_layout();v["pages"][0]["blocks"]=[{"id":"date","type":"text","x":40,"y":40,"width":600,"height":70,"text":"Referência: {{data_referencia}}"}]
    content,_=export(r,v,{"data_referencia":"2026-10-25"},[],"pdf")
    assert "25/10/2026" in PdfReader(io.BytesIO(base64.b64decode(content))).pages[0].extract_text()
    assert display(date(2026,10,25))=="25/10/2026"


def test_multi_series_pdf_has_readable_legend(env):
    _,r,_=env
    v=layout();v["sources"][0]["columns"].append("Meta");v["sources"][0]["rows"]=[["Jan",10,20],["Fev",20,30]]
    v["pages"][0]["blocks"][2]["options"]["y_fields"]=["Total","Meta"]
    content,_=export(r,v,{},[],"pdf")
    text=PdfReader(io.BytesIO(base64.b64decode(content))).pages[0].extract_text()
    assert "Total" in text and "Meta" in text


def test_canvas_accepts_legacy_shorthand_parameters(env):
    e,r,_=env
    r.source="local";r.handler="contacts";r.parameter_schema={"start":"date","end":"date"};r.save()
    response=post(e,r,"canvas-preview",{"layout":layout(),"parameters":{"start":"2026-10-01","end":"2026-10-25"},"output_format":"pdf"})
    assert response.status_code==200,response.data
    job=make_execution(r,{"start":"2026-10-01","end":"2026-10-25"},user=e["user"])
    execute_report(job.pk);job.refresh_from_db();assert job.state=="ready",job.error_code


def test_definition_signs_repeated_source_once_but_rechecks_next_snapshot(env):
    from operations.report_canvas.sources import signature
    _,report,child=env
    report.canvas_layout["sources"] += [{"id":"a","type":"report","report_id":child.pk},{"id":"b","type":"report","report_id":child.pk}]
    with patch("operations.report_canvas.engine.signature",wraps=signature) as sign:
        first=definition_snapshot(report)
        assert sign.call_count==1
        second=definition_snapshot(report)
        assert sign.call_count==2 and first==second
        child.version+=1;child.save()
        changed=definition_snapshot(report)
        assert sign.call_count==3 and changed!=first


def test_repeated_source_still_validates_each_dataset_reference(env):
    _,report,child=env
    report.canvas_layout["sources"] += [{"id":"a","type":"report","report_id":child.pk},{"id":"b","type":"report","report_id":child.pk,"dataset":"not_allowed"}]
    with pytest.raises(ValidationError):definition_snapshot(report)


def test_snapshot_asset_metadata_does_not_load_image_bytes(env):
    from operations.models import ReportCanvasAsset
    from operations.report_canvas.engine import assets_for
    _,report,_=env
    image=ReportCanvasAsset.objects.create(unit=report.unit,report=report,filename="test.png",width=1,height=1,data=b"synthetic",sha256="0"*64)
    report.canvas_layout["pages"][0]["blocks"].append({"id":"image","type":"image","x":1,"y":1,"width":20,"height":20,"asset_id":str(image.pk)})
    assert "data" in assets_for(report,report.canvas_layout,metadata_only=True)[str(image.pk)].get_deferred_fields()
    assert "data" not in assets_for(report,report.canvas_layout)[str(image.pk)].get_deferred_fields()
