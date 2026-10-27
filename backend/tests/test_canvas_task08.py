"""Task 08: isolated adversarial tests; never use the staging database."""
import base64
import io
import json
import math
import uuid
from types import SimpleNamespace
from unittest.mock import patch
import pytest
from rest_framework.exceptions import ValidationError
from core.models import Audit, Membership, Unit
from operations.models import ReportDefinition
from operations.report_canvas.api import serialize_cell
from operations.report_canvas.blocks import BLOCKS, aggregate_series, parse_group_date
from operations.report_canvas.schema import default_layout, validate_layout
from operations.report_canvas.sources import Dataset, InlineSource


def chart(rows, kind="bar", **options):
    b = {"id": "c", "type": kind, "source_id": "s", "options": {"x_field": "X", "y_fields": ["Y"], **options}}
    return BLOCKS[kind].prepare(b, SimpleNamespace(datasets={"s": Dataset(["X", "Y"], rows)}))


@pytest.mark.parametrize("count", [0, 1, 500])
def test_chart_boundaries(count):
    assert len(chart([[str(i), 0] for i in range(count)])["rows"]) == count


def test_chart_501_points_rejected():
    with pytest.raises(ValidationError) as exc:
        chart([[str(i), 1] for i in range(501)])
    assert exc.value.detail["code"] == "chart_points_limit"


@pytest.mark.parametrize("value", [math.nan, math.inf, -math.inf, "NaN", "Infinity"])
def test_nonfinite_series(value):
    with pytest.raises(ValidationError):
        chart([["A", value]])


@pytest.mark.parametrize("value", [math.nan, math.inf, -math.inf])
def test_inline_nonfinite_rejected(value):
    with pytest.raises(ValidationError):
        InlineSource().validate({"columns": ["X"], "rows": [[value]]})


def test_null_zero_and_negative_pie():
    assert chart([["A", None], ["B", 0]])["rows"] == [["A", None], ["B", 0.0]]
    with pytest.raises(ValidationError) as exc:
        chart([["A", -1]], "pie")
    assert exc.value.detail["code"] == "pie_negative_value"


def test_aggregate_overflow_is_actionable():
    with pytest.raises(ValidationError) as exc:
        aggregate_series("sum", [1e308, 1e308], "Y", {"id": "c"})
    assert exc.value.detail["code"] == "series_not_numeric"
    assert aggregate_series("avg", [1e308, 1e308], "Y", {"id": "c"}) == 1e308


@pytest.mark.parametrize("value", ["2023-02-29", "2024-02-30", "2026-13-01", "2026-01-01T25:00:00Z", "2026-01-01T12:00:00+25:00", "20261027", "2020-W53-4", "pedido 2026-01-01"])
def test_invalid_dates(value):
    with pytest.raises(ValidationError) as exc:
        parse_group_date(value, {"id": "c"}, "X")
    assert exc.value.detail["code"] == "x_not_a_date"


def test_leap_timezone_and_week53():
    assert chart([[None, 99], ["2024-02-29T23:30:00-03:00", 1], ["2024-02-29T01:00:00+04:00", 2]], x_group="day")["rows"] == [["29/02/2024", 3.0]]
    assert chart([["2020-12-31", 1], ["2021-01-01", 2], ["2021-01-04", 4]], x_group="week")["rows"] == [["sem. 28/12", 3.0], ["sem. 04/01", 4.0]]


def text_block(i, text="X"):
    return {"id": f"b{i}", "type": "text", "x": 0, "y": 0, "width": 20, "height": 20, "text": text}


def test_layout_limits():
    v = default_layout()
    v["page"].update(width=150, height=150)
    v["sources"] = [{"id": f"s{i}", "type": "goals"} for i in range(30)]
    v["pages"] = [{"id": f"p{i}", "name": "Página", "blocks": []} for i in range(30)]
    v["pages"][0]["blocks"] = [text_block(i) for i in range(100)]
    assert validate_layout(v)
    v["page"].update(width=3000, height=3000)
    assert validate_layout(v)
    for key, item in [("sources", {"id": "extra", "type": "goals"}), ("pages", {"id": "extra", "name": "Extra", "blocks": []})]:
        v[key].append(item)
        with pytest.raises(ValidationError): validate_layout(v)
        v[key].pop()
    v["pages"][0]["blocks"].append(text_block(100))
    with pytest.raises(ValidationError): validate_layout(v)


def test_layout_size_and_long_text_limits():
    v = default_layout()
    v["pages"][0]["blocks"] = [text_block(i, "a" * 20000) for i in range(36)]
    assert 720000 < len(json.dumps(v).encode()) < 750000
    assert validate_layout(v)
    v["pages"][0]["blocks"].append(text_block(36, "a" * 20000))
    v["pages"][0]["blocks"].append(text_block(37, "a" * 20000))
    with pytest.raises(ValidationError, match="750"):
        validate_layout(v)
    v["pages"][0]["blocks"] = [text_block(0, "a" * 20001)]
    with pytest.raises(ValidationError): validate_layout(v)


def test_manual_table_limits_and_names():
    cols = ["á" * 100, "Valor total"] + [f"C{i}" for i in range(98)]
    source = {"columns": cols, "rows": [[None] * 100 for _ in range(2000)]}
    InlineSource().validate(source)
    source["rows"].append([None] * 100)
    with pytest.raises(ValidationError): InlineSource().validate(source)
    with pytest.raises(ValidationError): InlineSource().validate({"columns": ["á", "á"], "rows": []})


@pytest.fixture
def canvas_env(setup):
    report = ReportDefinition.objects.create(unit=setup["unit"], code="task08", name="QA", source="legacy_pending", enabled=True, reviewed=True)
    return setup, report, f"/api/v1/reports/{report.pk}/canvas-source-preview/?unit=ame"


def test_preview_permissions_and_deactivated_source(canvas_env):
    env, report, url = canvas_env
    foreign = Unit.objects.create(name="Outra", slug="other-qa")
    child = ReportDefinition.objects.create(unit=foreign, code="foreign", name="Outra", source="legacy_pending", reviewed=True, enabled=True)
    payload = {"source": {"id": "s", "type": "report", "report_id": child.pk}}
    assert env["client"].post(url, payload, format="json").status_code == 400
    assert env["client"].post(url.replace(str(report.pk), str(child.pk), 1), payload, format="json").status_code == 404
    child.unit = env["unit"]; child.enabled = False; child.save()
    assert env["client"].post(url, payload, format="json").status_code == 400
    Membership.objects.filter(user=env["user"]).update(role="operator")
    assert env["client"].post(url, {"source": {"id": "s", "type": "goals"}}, format="json").status_code == 403


def test_preview_exception_does_not_disclose_connection_or_data(canvas_env, caplog):
    env, _, url = canvas_env
    marker = "PRIVATE_SQL_DATA_AND_CONNECTION"
    with patch("operations.report_canvas.sources.GoalsSource.read", side_effect=RuntimeError(marker)):
        response = env["client"].post(url, {"source": {"id": "s", "type": "goals"}}, format="json")
    assert response.status_code == 400
    assert response.data["code"] == "source_connection_failed"
    assert marker not in str(response.data)
    assert marker not in caplog.text


def test_preview_audit_contains_no_sample_data(canvas_env):
    env, report, url = canvas_env
    marker = "PRIVATE_SAMPLE_VALUE"
    response = env["client"].post(url, {"source": {"id": "s", "type": "inline", "columns": ["X"], "rows": [[marker]]}}, format="json")
    assert response.status_code == 200
    assert marker not in json.dumps(list(Audit.objects.filter(action="report.canvas_source_previewed", object_id=str(report.pk)).values_list("details", flat=True)))


@pytest.mark.parametrize("fmt", ["PNG", "JPEG"])
def test_assets_boundaries_and_corrupt_file(canvas_env, fmt):
    from PIL import Image
    env, report, _ = canvas_env
    url = f"/api/v1/reports/{report.pk}/canvas-assets/?unit=ame"
    def upload(raw):
        return env["client"].post(url, {"filename": "../ação especial.png", "content": base64.b64encode(raw).decode()}, format="json", HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()))
    out = io.BytesIO(); Image.new("RGB", (4000, 4000), "white").save(out, format=fmt)
    result = upload(out.getvalue())
    assert result.status_code == 200
    assert (result.data["width"], result.data["height"]) == (4000, 4000)
    assert result.data["filename"] == "ação especial.png"
    assert upload(b"corrupt image").status_code == 400
    assert upload(b"x" * (2 * 1024 * 1024 + 1)).status_code == 400
    out = io.BytesIO(); Image.new("RGB", (4001, 4000), "white").save(out, format=fmt)
    assert upload(out.getvalue()).status_code == 400


@pytest.mark.parametrize("path,value", [
    (("xlsx_mode",), []), (("sources", 0, "type"), []),
    (("pages", 0, "blocks", 0, "type"), {}),
    (("pages", 0, "blocks", 0, "source_id"), ["s"]),
    (("pages", 0, "blocks", 0, "options", "aggregate"), {}),
    (("pages", 0, "blocks", 0, "x"), 10 ** 1000),
])
def test_malformed_json_never_raises_server_error(path, value):
    v = default_layout()
    v["sources"] = [{"id": "s", "type": "goals"}]
    v["pages"][0]["blocks"] = [{**text_block(1), "options": {}}]
    target = v
    for key in path[:-1]: target = target[key]
    target[path[-1]] = value
    with pytest.raises(ValidationError): validate_layout(v)


@pytest.mark.parametrize("source", [
    {"id": "s", "type": []}, {"id": "s", "type": {}},
    {"id": "s", "type": "goals", "label": "a" * 61},
    {"id": ["s"], "type": "goals"},
])
def test_source_preview_malformed_json_is_400(canvas_env, source):
    env, _, url = canvas_env
    assert env["client"].post(url, {"source": source}, format="json").status_code == 400


def test_manual_table_combined_layout_limit():
    v = default_layout()
    v["sources"] = [{"id": "s", "type": "inline", "columns": [f"C{i}" for i in range(100)], "rows": [[0] * 100 for _ in range(2000)]}]
    assert validate_layout(v)


def test_20_image_limit_and_exact_2mb_png(canvas_env):
    from PIL import Image
    from operations.models import ReportCanvasAsset
    env, report, _ = canvas_env
    url = f"/api/v1/reports/{report.pk}/canvas-assets/?unit=ame"
    raw = io.BytesIO(); Image.new("RGB", (1, 1), "white").save(raw, format="PNG")
    payload = raw.getvalue().ljust(2 * 1024 * 1024, b"\0")
    result = env["client"].post(url, {"filename": "ação.png", "content": base64.b64encode(payload).decode()}, format="json", HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()))
    assert result.status_code == 200
    ReportCanvasAsset.objects.bulk_create([ReportCanvasAsset(unit=env["unit"], report=report, filename="qa.png", data=raw.getvalue(), width=1, height=1, sha256="0" * 64) for _ in range(19)])
    result = env["client"].post(url, {"filename": "qa.png", "content": base64.b64encode(raw.getvalue()).decode()}, format="json", HTTP_IDEMPOTENCY_KEY=str(uuid.uuid4()))
    assert result.status_code == 400
    assert "20 imagens" in str(result.data)


def test_literal_html_unknown_variable_and_newlines_export(canvas_env):
    from operations.report_canvas.engine import export
    from openpyxl import load_workbook
    from pypdf import PdfReader
    _, report, _ = canvas_env
    v = default_layout()
    text = "<script>alert(1)</script>\n{{inexistente}}\nSegunda linha"
    v["pages"][0]["blocks"] = [{**text_block(1, text), "width": 700, "height": 300}]
    encoded, _ = export(report, v, {}, [], "pdf")
    pdf_text = PdfReader(io.BytesIO(base64.b64decode(encoded))).pages[0].extract_text()
    assert "<script>alert(1)</script>" in pdf_text
    assert "{{inexistente}}" in pdf_text
    assert "Segunda linha" in pdf_text
    v["pages"][0]["blocks"][0]["text"] += " 😀"
    encoded, _ = export(report, v, {}, [], "xlsx")
    book = load_workbook(io.BytesIO(base64.b64decode(encoded)))
    assert book.active.cell(1, 1).value == text + " 😀"
    assert book.active.cell(1, 1).data_type == "s"


@pytest.mark.parametrize("value", [math.nan, math.inf, -math.inf, __import__("decimal").Decimal("NaN"), __import__("decimal").Decimal("Infinity"), __import__("decimal").Decimal("1.1e999")])
def test_source_preview_nonfinite_serialization_is_actionable(value):
    with pytest.raises(ValidationError) as exc: serialize_cell(value)
    assert exc.value.detail["code"] == "source_nonfinite_value"
