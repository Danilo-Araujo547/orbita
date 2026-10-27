import io
import math
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter
from .registry import EXPORTERS, Registry

RENDERERS=Registry()


def cell(sheet,row,column,value):
    target=sheet.cell(row,column,value)
    if isinstance(value,str):target.data_type="s"
    from datetime import date, datetime
    if isinstance(value,datetime):target.number_format="dd/mm/yyyy hh:mm:ss"
    elif isinstance(value,date):target.number_format="dd/mm/yyyy"
    return target


def dataset_sheet(book,b,prefix):
    title=re.sub(r"[\\/*?:\[\]]"," ",b.get("title") or b["id"])
    name=f"{prefix}{len(book.worksheets)} {title}"[:31]
    sheet=book.create_sheet(name)
    for ri,row in enumerate([b["columns"],*b["rows"]],1):
        for ci,value in enumerate(row,1):cell(sheet,ri,ci,value)
    sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
    for c in sheet[1]:c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="123C57")
    weights=b["options"].get("column_widths")
    for col in range(1,len(b["columns"])+1):sheet.column_dimensions[get_column_letter(col)].width=min(80,max(8,weights[col-1]/8)) if weights else 22
    return sheet


def text(ctx,b,row,col,width,height):
    cell(ctx.sheet,row,col,(b.get("title","")+"\n"+b.get("text","")).strip()).alignment=Alignment(wrap_text=True,vertical="top")
    target=ctx.sheet.cell(row,col);target.font=Font(size=b["options"].get("font_size",14),color=b["options"].get("color","#075b94")[1:])
    if width>1:ctx.sheet.merge_cells(start_row=row,start_column=col,end_row=row,end_column=col+width-1)
    ctx.sheet.row_dimensions[row].height=max(20,min(300,b["height"]*.75))


def kpi(ctx,b,row,col,width,height):
    cell(ctx.sheet,row,col,b.get("title","")).font=Font(bold=True,size=11)
    c=cell(ctx.sheet,row+1,col,b["value"] if b["value"] is not None else "Sem dados");c.font=Font(bold=True,size=24,color=b["options"].get("color","#075b94")[1:]);c.number_format="#,##0"+("."+"0"*b["options"].get("precision",0) if b["options"].get("precision",0) else "")


def image(ctx,b,row,col,width,height):
    img=Image(io.BytesIO(bytes(b["image"].data)))
    ratio=min(b["width"]/img.width,b["height"]/img.height);img.width*=ratio;img.height*=ratio
    ctx.sheet.add_image(img,f"{get_column_letter(col)}{row}")


def table(ctx,b,row,col,width,height):
    has_title=bool(b.get("title"))
    if has_title:cell(ctx.sheet,row,col,b["title"]).font=Font(bold=True)
    start=row+int(has_title)
    available=max(1,height-int(has_title)-1)
    visible_columns=max(1,width)
    for ri,line in enumerate([b["columns"][:visible_columns],*[r[:visible_columns] for r in b["rows"][:available]]]):
        for ci,value in enumerate(line):
            c=cell(ctx.sheet,start+ri,col+ci,value);c.alignment=Alignment(wrap_text=True,vertical="top")
            if ri==0:c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="123C57")
    if len(b["rows"])>available or len(b["columns"])>visible_columns:
        data=dataset_sheet(ctx.book,b,"Dados ")
        c=cell(ctx.sheet,start+available+1,col,"Dados completos");c.hyperlink=f"#'{data.title}'!A1";c.style="Hyperlink"


def chart(ctx,b,row,col,width,height):
    data=dataset_sheet(ctx.book,b,"_Série ");data.sheet_state="hidden"
    g={"bar":BarChart,"line":LineChart,"pie":PieChart}[b["type"]]()
    g.title=b.get("title","");g.style=13;g.width=b["width"]/37.795;g.height=b["height"]/37.795
    if b["rows"]:
        g.add_data(Reference(data,min_col=2,max_col=2 if b["type"]=="pie" else len(b["columns"]),min_row=1,max_row=len(b["rows"])+1),titles_from_data=True)
        g.set_categories(Reference(data,min_col=1,min_row=2,max_row=len(b["rows"])+1))
        ctx.sheet.add_chart(g,f"{get_column_letter(col)}{row}")
    else:cell(ctx.sheet,row,col,"Sem dados")


class XLSXExporter:
    def render(self,document):
        from types import SimpleNamespace
        book=Workbook();book.remove(book.active)
        if document.get("xlsx_mode") == "table":
            b=next(b for p in document["pages"] for b in p["blocks"] if b["type"]=="table")
            sheet=dataset_sheet(book,b,"")
            previous=object();band=False
            for ri,line in enumerate(b["rows"],2):
                key=line[-1]
                changed=key!=previous
                if changed:band=not band
                for c in sheet[ri]:
                    c.fill=PatternFill("solid",fgColor="EDF5FA" if band else "FFFFFF")
                    c.alignment=Alignment(vertical="center")
                    if changed:c.border=Border(top=Side(style="thin",color="9BB7CA"))
                previous=key
            out=io.BytesIO();book.save(out);return out.getvalue()
        page=document["page"];unit=page["width"]/page["columns"]
        for index,p in enumerate(document["pages"],1):
            name=f"{index} "+re.sub(r"[\\/*?:\[\]]"," ",p["name"])
            sheet=book.create_sheet(name[:31]);sheet.sheet_view.showGridLines=False
            for i in range(1,page["columns"]+1):sheet.column_dimensions[get_column_letter(i)].width=max(3,unit/7)
            sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation="landscape" if page["width"]>page["height"] else "portrait"
            sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=1
            ctx=SimpleNamespace(book=book,sheet=sheet)
            for b in p["blocks"]:
                row=int(b["y"]//24)+1;col=int(b["x"]//unit)+1;width=max(1,int(b["width"]//unit));height=max(1,int(b["height"]//24))
                RENDERERS[b["type"]](ctx,b,row,col,width,height)
            sheet.print_options.horizontalCentered=True
        out=io.BytesIO();book.save(out);return out.getvalue()

RENDERERS.register("text",text)
RENDERERS.register("kpi",kpi)
RENDERERS.register("image",image)
RENDERERS.register("table",table)
for kind in ("line","bar","pie"):RENDERERS.register(kind,chart)
EXPORTERS.register("xlsx",XLSXExporter())
