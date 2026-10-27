import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime
from xml.etree.ElementTree import ParseError
from zoneinfo import ZoneInfo

from django.db import transaction
from rest_framework.exceptions import ValidationError

from operations.absence_html import parse_absence
from operations.demand_csv import parse_demand
from operations.models import Case, Contact, SheetImport, SourceEvent
from operations.services import log
from operations.standardization import (
    DEFAULT_MAPPINGS,
    standard_date,
    standard_row,
    template_bindings,
    unique_headers,
    validate_mapping,
)
from scheduling.connectors import normalize_phone
from scheduling.models import Patient, PatientPhone

REQUIRED = {"external_id", "patient", "phone", "occurred_at"}


def preview_sheet(
    unit,
    layout,
    actor,
    filename,
    content,
    *,
    template_id=None,
    standard_mapping=None,
    parameter_fields=None,
    source_rows=None,
    source_metadata=None,
):

    if layout.unit_id != unit.id or not layout.verified:
        raise ValidationError("Layout não validado para esta unidade.")

    if source_rows is None and len(content) > 5 * 1024 * 1024:
        raise ValidationError("Arquivo excede 5 MB.")

    if layout.format == "table" and not REQUIRED.issubset(layout.mapping):
        raise ValidationError(
            "Layout precisa mapear identificador, paciente, telefone e data."
        )

    from messaging.models import Template

    template = None
    selected_id = template_id or layout.default_template_id
    if selected_id:
        template = Template.objects.filter(pk=selected_id, unit=unit).first()
        if not template:
            raise ValidationError("Template não pertence à unidade.")
        allowed_purposes = (
            {"absence"}
            if layout.format == "absence_html"
            else {"waitlist"}
            if layout.format == "cross_demand_csv"
            else {"waitlist", "summons"}
        )
        if template.purpose not in allowed_purposes:
            raise ValidationError(
                "Finalidade do template incompatível com o modelo da planilha."
            )
    mapping = validate_mapping(
        standard_mapping
        if standard_mapping is not None
        else layout.standard_mapping or DEFAULT_MAPPINGS[layout.format]
    )
    fields = (
        template_bindings(
            template,
            parameter_fields
            if parameter_fields is not None
            else layout.parameter_fields or None,
        )
        if template
        else []
    )
    configuration = {
        "schema_version": 1,
        "layout_version": layout.version,
        "standard_mapping": mapping,
        "parameter_fields": fields,
    }
    if source_metadata is not None:
        configuration["source"] = source_metadata
    if template:
        configuration["template"] = {
            "id": template.pk,
            "version": template.version,
            "name": template.name,
            "language": template.language,
            "purpose": template.purpose,
            "body": template.body,
            "components": template.components,
        }
    digest = hashlib.sha256(
        content
        + (
            json.dumps(configuration, sort_keys=True).encode()
            if template or standard_mapping is not None or source_metadata is not None
            else b""
        )
    ).hexdigest()

    existing = SheetImport.objects.filter(
        unit=unit, layout=layout, digest=digest
    ).first()

    if existing:
        return existing

    try:
        if source_rows is not None:
            source = source_rows
        elif layout.format == "absence_html":
            if not filename.lower().endswith(".xls"):
                raise ValueError("unsupported_file")

            source = parse_absence(content)

        elif layout.format == "cross_demand_csv":
            if not filename.lower().endswith(".csv"):
                raise ValueError("unsupported_file")
            source = parse_demand(content)

        elif filename.lower().endswith(".csv"):
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp1252")

            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")

            reader = csv.reader(io.StringIO(text), dialect=dialect)
            headers = unique_headers(next(reader))
            source = (dict(zip(headers, values)) for values in reader)

        elif filename.lower().endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError("expanded_file_too_large")

            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )

            iterator = workbook.active.iter_rows(values_only=True)

            headers = unique_headers(next(iterator))

            source = (dict(zip(headers, values)) for values in iterator)

        else:
            raise ValueError("unsupported_file")

        rows, errors, seen, duplicates = [], [], {}, 0

        for number, raw in enumerate(source, 2):
            if number > 10001:
                raise ValueError("row_limit")

            row = (
                dict(raw)
                if layout.format in {"absence_html", "cross_demand_csv"}
                else {
                    field: str(raw.get(column) or "").strip()
                    for field, column in layout.mapping.items()
                }
            )

            try:
                if row.get("_error"):
                    raise ValueError(row["_error"])

                if layout.format == "absence_html" and (
                    not row.get("patient_external_id")
                    or len(row["patient_external_id"]) > 100
                ):
                    raise ValueError("missing_fields")

                if layout.format == "table" and (
                    standard_mapping is not None or layout.standard_mapping
                ):
                    standard_input = standard_row(raw, {}, mapping, "table")
                    if mapping.get("codigo"):
                        row["external_id"] = standard_input["codigo"]
                    if mapping.get("nome_completo"):
                        row["patient"] = standard_input["nome_completo"]
                    if mapping.get("celular") or mapping.get("celular2"):
                        row["phone"] = (
                            standard_input["celular"] or standard_input["celular2"]
                        )
                    reference = str(
                        raw.get(layout.mapping.get("occurred_at", "")) or ""
                    ).strip()
                    row["occurred_at"] = reference or standard_input["data"]

                if any(not row.get(field) for field in REQUIRED):
                    raise ValueError("missing_fields")

                if len(row["external_id"]) > 200 or len(row["patient"]) > 200:
                    raise ValueError("field_too_long")

                row["phone"] = normalize_phone(row["phone"])

                try:
                    day = datetime.fromisoformat(row["occurred_at"])
                except ValueError:
                    day = datetime.strptime(
                        standard_date(row["occurred_at"]), "%d/%m/%Y"
                    )
                    time_value = str(raw.get(mapping.get("hora", "")) or "").strip()
                    if time_value:
                        from operations.standardization import standard_time

                        hour = datetime.strptime(
                            standard_time(time_value), "%H:%M"
                        ).time()
                        day = datetime.combine(day.date(), hour)

                if day.tzinfo is None:
                    day = day.replace(tzinfo=ZoneInfo(unit.timezone))

                row["occurred_at"] = day.isoformat()
                standard = standard_row(raw, row, mapping, layout.format)
                row.update(standard)
                row["standardized"] = standard
                if standard["nome_completo"]:
                    row["patient"] = standard["nome_completo"]
                numbers = [
                    standard[key] for key in ["celular", "celular2"] if standard[key]
                ]
                source_numbers = row.get("phone_numbers")
                if not isinstance(source_numbers, list):
                    source_numbers = [row["phone"]]
                row["phone_numbers"] = list(dict.fromkeys(numbers + source_numbers))
                row["phone"] = row["phone_numbers"][0]
                if template and any(
                    not standard.get(field, "").strip() for field in fields
                ):
                    raise ValueError("missing_template_parameter")

                key = row["external_id"]

                if key in seen:
                    if seen[key] != row:
                        raise ValueError("conflicting_identifier")

                    duplicates += 1

                    continue

                seen[key] = row

                rows.append(row)

            except ValueError as exc:
                errors.append(
                    {
                        "line": number,
                        "code": str(exc)
                        if str(exc)
                        in {
                            "missing_fields",
                            "field_too_long",
                            "conflicting_identifier",
                            "invalid_phone_format",
                            "invalid_date_or_phone",
                            "missing_template_parameter",
                            "invalid_standard_date",
                            "invalid_standard_time",
                        }
                        else "invalid_date",
                    }
                )

        if not rows:
            errors.append({"line": 0, "code": "empty_file"})

    except (
        ValueError,
        UnicodeError,
        StopIteration,
        csv.Error,
        OSError,
        KeyError,
        zipfile.BadZipFile,
        ParseError,
        EOFError,
    ):
        raise ValidationError(
            "Arquivo inválido ou incompatível com o modelo selecionado."
        ) from None

    finally:
        if "workbook" in locals():
            workbook.close()

    return SheetImport.objects.create(
        unit=unit,
        layout=layout,
        digest=digest,
        rows=rows,
        errors=errors,
        duplicates=duplicates,
        created_by=actor,
        filename=filename,
        template=template,
        configuration=configuration,
    )


@transaction.atomic
def confirm_sheet(batch, purpose):

    batch = SheetImport.objects.select_for_update().get(pk=batch.pk)

    allowed = (
        {"absence"}
        if batch.layout.format == "absence_html"
        else {"waitlist"}
        if batch.layout.format == "cross_demand_csv"
        else {"waitlist", "summons"}
    )

    if purpose not in allowed or batch.errors or batch.state != "preview":
        raise ValidationError(
            "Importação precisa estar em prévia, sem erros, para CDR."
        )

    if batch.template_id:
        frozen = batch.configuration.get("template", {})
        if (
            batch.template.unit_id != batch.unit_id
            or batch.template.purpose != purpose
            or frozen.get("version") != batch.template.version
        ):
            raise ValidationError(
                "Template alterado ou incompatível com a finalidade da importação."
            )

    if not batch.layout.verified:
        raise ValidationError("Layout deixou de estar validado.")

    cases = []

    for row in batch.rows:
        event, created = SourceEvent.objects.get_or_create(
            unit=batch.unit,
            source=f"absence:{batch.layout_id}"
            if purpose == "absence"
            else f"cross:{batch.layout_id}",
            kind="absence" if purpose == "absence" else "waitlist",
            external_id=row["external_id"],
            defaults={
                "patient": row["patient"],
                "phone": row["phone"],
                "occurred_at": datetime.fromisoformat(row["occurred_at"]),
                "data": row,
                "synthetic": False,
            },
        )

        if not created and (
            event.patient != row["patient"]
            or event.phone != row["phone"]
            or (
                (
                    batch.template_id
                    or batch.layout.format in {"absence_html", "cross_demand_csv"}
                )
                and event.data != row
            )
        ):
            raise ValidationError(
                "Identificador já existente com contato ou paciente diferente; revisão necessária."
            )

        case, new = Case.objects.get_or_create(
            unit=batch.unit, event=event, purpose=purpose
        )

        if new:
            log(
                case,
                "imported",
                "Registro gerado pelo Oracle."
                if batch.configuration.get("source", {}).get("origin")
                == "oracle_absences"
                else "Caso criado por importação CROSS.",
                actor=batch.created_by,
                data={"batch": str(batch.pk), "layout_version": batch.layout.version},
            )

        if purpose == "absence":
            patient, _ = Patient.objects.get_or_create(
                unit=batch.unit,
                source=event.source,
                external_id=row["patient_external_id"],
                defaults={"name": row["patient"]},
            )

            if patient.name != row["patient"]:
                raise ValidationError(
                    "Código de paciente conflitante; revisão necessária."
                )

            for phone in row["phone_numbers"]:
                patient_phone, _ = PatientPhone.objects.get_or_create(
                    patient=patient, phone=phone
                )

                Contact.objects.get_or_create(
                    unit=batch.unit,
                    event=event,
                    phone=phone,
                    defaults={"patient_phone": patient_phone},
                )

        else:
            for phone in row.get("phone_numbers", [event.phone]):
                Contact.objects.get_or_create(unit=batch.unit, event=event, phone=phone)

        cases.append(str(case.pk))

    batch.state = "confirmed"

    batch.version += 1
    batch.save(update_fields=["state", "version"])
    if batch.template_id:
        from operations.models import Campaign

        Campaign.objects.create(
            unit=batch.unit,
            import_batch=batch,
            name=batch.filename[:150] or f"Importação {batch.pk}",
            purpose=purpose,
            template=batch.template,
        ).cases.set(cases)

    return cases
